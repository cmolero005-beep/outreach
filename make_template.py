"""Builds outreach_template.xlsx. Re-run only if you want to regenerate the blank template."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

HEADERS = [
    ("First Name", 14, True), ("Last Name", 14, True), ("Email", 30, True),
    ("Company Name", 20, True), ("Role", 20, True), ("Area", 16, True), ("Commonality", 32, True),
    ("Status", 13, False), ("Date Sent", 12, False), ("Delivery", 22, False),
    ("Replied", 9, False), ("Follow Up", 12, False),
]

wb = Workbook()
ws = wb.active
ws.title = "Contacts"
you = PatternFill("solid", fgColor="1F4E78")
auto = PatternFill("solid", fgColor="7F7F7F")
for i, (name, width, yours) in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = you if yours else auto
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[c.column_letter].width = width
for i, name in enumerate(["Message ID", "Subject Sent"], start=len(HEADERS) + 1):
    c = ws.cell(row=1, column=i, value=name)
    ws.column_dimensions[c.column_letter].hidden = True
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:L1"

status = DataValidation(type="list", formula1='"Sent,Followed Up,Replied,Bounced,Failed,Do Not Contact"',
                        allow_blank=True)
replied = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(status); status.add("H2:H1000")
ws.add_data_validation(replied); replied.add("K2:K1000")

green = PatternFill("solid", fgColor="C6EFCE"); red = PatternFill("solid", fgColor="FFC7CE")
blue = PatternFill("solid", fgColor="DDEBF7"); yellow = PatternFill("solid", fgColor="FFEB9C")
ws.conditional_formatting.add("A2:L1000", FormulaRule(formula=['$H2="Replied"'], fill=green))
ws.conditional_formatting.add("A2:L1000", FormulaRule(formula=['OR($H2="Bounced",$H2="Failed")'], fill=red))
ws.conditional_formatting.add("A2:L1000", FormulaRule(formula=['$H2="Followed Up"'], fill=yellow))
ws.conditional_formatting.add("A2:L1000", FormulaRule(formula=['$H2="Sent"'], fill=blue))

h = wb.create_sheet("How To Use")
h.column_dimensions["A"].width = 110
for r, line in enumerate([
    "HOW TO USE",
    "",
    "1. On the Contacts tab, fill in the DARK BLUE columns: First Name, Email, Area and Commonality are needed.",
    "   Last Name, Company Name and Role are optional (only needed if your email draft uses them).",
    "2. Leave Status EMPTY for new people. Empty Status = 'send the first email at the next 11:00 AM weekday run'.",
    "3. The GREY columns are filled in automatically - don't type in them:",
    "     Status      Sent -> Followed Up -> Replied   (or Bounced / Failed)",
    "     Date Sent   the day the first email went out",
    "     Delivery    'Sent', 'Bounced', 'Failed: ...', or 'Waiting: fill in ...' if a field is missing",
    "     Replied     Yes / No - checked against your inbox (Mail app) every run",
    "     Follow Up   the date the follow-up was sent (only if they haven't replied within 7 days)",
    "4. To stop emailing someone, set their Status to 'Do Not Contact'. If they replied outside email, set Replied = Yes.",
    "5. CLOSE this file in Excel before 11:00 AM so the script can save its updates.",
    "   (If it's open, the script saves a copy called outreach_UPDATED_<date>.xlsx instead.)",
    "",
    "Row colours: blue = sent, yellow = followed up, green = replied, red = bounced/failed.",
], start=1):
    h.cell(row=r, column=1, value=line).font = Font(bold=(r == 1), size=14 if r == 1 else 11)

wb.save("outreach_template.xlsx")
print("wrote outreach_template.xlsx")
