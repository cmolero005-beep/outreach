#!/usr/bin/env python3
"""
Networking email outreach automation.

Reads contacts from an Excel sheet, sends a personalized first email to every
new row, sends one follow-up to people who haven't replied after N days
(7 by default), and keeps the Status / Date Sent / Delivery / Replied / Follow Up columns
up to date.

Commands:
    python outreach.py init        create outreach.xlsx from the template
    python outreach.py preview     show the emails that WOULD be sent (sends nothing)
    python outreach.py check       make sure the script can talk to your mail app
    python outreach.py test        send both emails to yourself using a sample contact
    python outreach.py run         the daily job (the scheduler calls this at 11:00)
    python outreach.py run --force   run even on a weekend
"""

import argparse
import configparser
import datetime as dt
import imaplib
import logging
import mimetypes
import random
import re
import shutil
import smtplib
import sys
import time
from email.message import EmailMessage
from email.utils import formataddr, make_msgid
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "config.ini"
TEMPLATE_XLSX = HERE / "outreach_template.xlsx"
LOG_DIR = HERE / "logs"

# Column headers in the sheet (order doesn't matter; matched by name).
COL_FIRST = "First Name"
COL_LAST = "Last Name"
COL_EMAIL = "Email"
COL_COMPANY = "Company Name"
COL_ROLE = "Role"
COL_AREA = "Area"
COL_COMMON = "Commonality"
COL_STATUS = "Status"
COL_DATE_SENT = "Date Sent"
COL_DELIVERY = "Delivery"
COL_REPLIED = "Replied"
COL_FOLLOW_UP = "Follow Up"
COL_MSG_ID = "Message ID"  # hidden helper column, used to thread the follow-up
COL_SUBJECT = "Subject Sent"  # hidden helper column

REQUIRED_FOR_SEND = [COL_FIRST, COL_EMAIL, COL_AREA, COL_COMMON]

STATUS_SENT = "Sent"
STATUS_FOLLOWED_UP = "Followed Up"
STATUS_REPLIED = "Replied"
STATUS_BOUNCED = "Bounced"
STATUS_FAILED = "Failed"
STATUS_SKIP = {"do not contact", "skip", "paused"}

# [Placeholder] in the template -> sheet column
PLACEHOLDERS = {
    "first name": COL_FIRST,
    "firstname": COL_FIRST,
    "name": COL_FIRST,
    "last name": COL_LAST,
    "lastname": COL_LAST,
    "area": COL_AREA,
    "commonality": COL_COMMON,
    "company": COL_COMPANY,
    "company name": COL_COMPANY,
    "role": COL_ROLE,
}

log = logging.getLogger("outreach")


# --------------------------------------------------------------------------- config

def load_config():
    if not CONFIG_PATH.exists():
        sys.exit(f"Missing {CONFIG_PATH.name}. Copy config.example.ini to config.ini and fill it in.")
    cfg = configparser.ConfigParser()
    cfg.read(CONFIG_PATH, encoding="utf-8")
    return cfg


def setup_logging():
    LOG_DIR.mkdir(exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(LOG_DIR / "outreach.log", encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )


# --------------------------------------------------------------------------- templates

def load_template(path):
    """Template file: first line 'Subject: ...', blank line, then the body."""
    text = Path(path).read_text(encoding="utf-8").replace("\r\n", "\n")
    subject = None
    lines = text.split("\n")
    if lines and lines[0].lower().startswith("subject:"):
        subject = lines[0].split(":", 1)[1].strip()
        lines = lines[1:]
        while lines and not lines[0].strip():
            lines = lines[1:]
    body = "\n".join(lines).rstrip() + "\n"
    check_placeholders((subject or "") + body, path)
    return subject, body


def check_placeholders(text, path):
    for m in re.finditer(r"\[([^\]]+)\]", text):
        if m.group(1).strip().lower() not in PLACEHOLDERS:
            sys.exit(
                f"Unknown placeholder [{m.group(1)}] in {path}. "
                f"Allowed: {', '.join('[' + k.title() + ']' for k in PLACEHOLDERS)}"
            )


ATTACHMENT_TYPES = {".pdf", ".docx", ".doc"}


class TemplateSet:
    def __init__(self, folder):
        self.name = folder.name
        self.subj1, self.body1 = load_template(folder / "initial_email.txt")
        if not self.subj1:
            sys.exit(f"{folder.name}/initial_email.txt needs a 'Subject: ...' first line.")
        self.subj2, self.body2 = load_template(folder / "follow_up_email.txt")
        used = placeholders_used(self.subj1 + self.body1 + (self.subj2 or "") + self.body2)
        self.needed = REQUIRED_FOR_SEND + sorted(used - set(REQUIRED_FOR_SEND))
        # Any PDF/Word file in the folder (e.g. your resume) is attached to the FIRST email.
        self.attachments = sorted(f for f in folder.iterdir()
                                  if f.suffix.lower() in ATTACHMENT_TYPES and not f.name.startswith("~$"))
        self.missing_attachment = "attach" in self.body1.lower() and not self.attachments

    def first(self, row):
        return render(self.subj1, row), render(self.body1, row)

    def follow_up(self, row):
        orig = str(row.get(COL_SUBJECT) or render(self.subj1, row))
        subj = render(self.subj2, row) if self.subj2 else ("Re: " + orig if not orig.startswith("Re:") else orig)
        return subj, render(self.body2, row)


class Templates:
    """One folder per commonality: templates/<Commonality>/initial_email.txt + follow_up_email.txt.
    The row's Commonality value (e.g. SNHU) picks the folder, ignoring upper/lower case."""

    def __init__(self, cfg):
        root = HERE / cfg["settings"].get("templates_dir", "templates")
        self.sets = {d.name.strip().lower(): TemplateSet(d) for d in sorted(root.iterdir())
                     if d.is_dir() and (d / "initial_email.txt").exists()}
        if not self.sets:
            sys.exit(f"No templates found in {root}. Expected {root}/<Commonality>/initial_email.txt")

    def for_row(self, row):
        return self.sets.get(str(row.get(COL_COMMON) or "").strip().lower())

    def names(self):
        return ", ".join(t.name for t in self.sets.values())


def placeholders_used(text):
    return {PLACEHOLDERS[m.group(1).strip().lower()] for m in re.finditer(r"\[([^\]]+)\]", text or "")}


def render(text, row):
    def sub(m):
        return str(row.get(PLACEHOLDERS[m.group(1).strip().lower()]) or "").strip()
    return re.sub(r"\[([^\]]+)\]", sub, text or "")


# --------------------------------------------------------------------------- workbook

class Sheet:
    def __init__(self, path):
        self.path = Path(path)
        self.wb = load_workbook(self.path)
        self.ws = self.wb["Contacts"] if "Contacts" in self.wb.sheetnames else self.wb.active
        self.cols = {}
        for cell in self.ws[1]:
            if cell.value:
                self.cols[str(cell.value).strip()] = cell.column
        missing = [c for c in (COL_FIRST, COL_EMAIL, COL_AREA, COL_STATUS, COL_DATE_SENT,
                               COL_DELIVERY, COL_REPLIED, COL_FOLLOW_UP) if c not in self.cols]
        if missing:
            sys.exit(f"Sheet is missing columns: {missing}")
        # Add hidden helper columns if the sheet doesn't have them yet.
        for name in (COL_MSG_ID, COL_SUBJECT):
            if name not in self.cols:
                col = self.ws.max_column + 1
                self.ws.cell(row=1, column=col, value=name)
                self.ws.column_dimensions[self.ws.cell(row=1, column=col).column_letter].hidden = True
                self.cols[name] = col

    def rows(self):
        for r in range(2, self.ws.max_row + 1):
            row = {name: self.ws.cell(row=r, column=c).value for name, c in self.cols.items()}
            if not any(row.get(c) for c in (COL_FIRST, COL_EMAIL)):
                continue
            row["_row"] = r
            yield row

    def set(self, row, col, value, fmt=None):
        cell = self.ws.cell(row=row["_row"], column=self.cols[col], value=value)
        if fmt:
            cell.number_format = fmt
        row[col] = value

    def save(self):
        try:
            self.wb.save(self.path)
        except PermissionError:
            # Excel has the file open (Windows locks it). Don't lose the updates.
            alt = self.path.with_name(self.path.stem + f"_UPDATED_{dt.datetime.now():%Y%m%d_%H%M}.xlsx")
            self.wb.save(alt)
            log.error("Could not save %s (is it open in Excel?). Saved updates to %s instead — "
                      "close Excel and replace the original with this file.", self.path.name, alt.name)


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


# --------------------------------------------------------------------------- mail backends

class SmtpImapMailer:
    """Works with Gmail / Google Workspace (app password) and any provider that allows SMTP+IMAP."""

    def __init__(self, cfg):
        s = cfg["email"]
        self.address = s["address"].strip()
        self.name = s.get("display_name", "").strip()
        self.password = s["password"].strip()
        self.smtp_host = s["smtp_host"].strip()
        self.smtp_port = s.getint("smtp_port", 587)
        self.imap_host = s.get("imap_host", "").strip()
        self.smtp = None

    def _smtp(self):
        if self.smtp is None:
            if self.smtp_port == 465:
                self.smtp = smtplib.SMTP_SSL(self.smtp_host, self.smtp_port, timeout=60)
            else:
                self.smtp = smtplib.SMTP(self.smtp_host, self.smtp_port, timeout=60)
                self.smtp.starttls()
            self.smtp.login(self.address, self.password)
        return self.smtp

    def send(self, to, subject, body, in_reply_to=None, attachments=()):
        msg = EmailMessage()
        msg["From"] = formataddr((self.name, self.address)) if self.name else self.address
        msg["To"] = to
        msg["Subject"] = subject
        msg["Message-ID"] = make_msgid(domain=self.address.split("@")[-1])
        if in_reply_to:
            msg["In-Reply-To"] = in_reply_to
            msg["References"] = in_reply_to
        msg.set_content(body)
        for path in attachments:
            ctype = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
            maintype, subtype = ctype.split("/", 1)
            msg.add_attachment(Path(path).read_bytes(), maintype=maintype, subtype=subtype, filename=Path(path).name)
        self._smtp().send_message(msg)
        return msg["Message-ID"]

    def check_connection(self):
        self._smtp()
        if self.imap_host:
            self._imap().logout()

    def _imap(self):
        imap = imaplib.IMAP4_SSL(self.imap_host)
        imap.login(self.address, self.password)
        return imap

    def find_replies_and_bounces(self, contacts):
        """contacts: list of (email, since_date). Returns (replied: {email: date}, bounced: set)."""
        replied, bounced = {}, set()
        if not self.imap_host or not contacts:
            return replied, bounced
        imap = self._imap()
        try:
            imap.select("INBOX", readonly=True)
            for email_addr, since in contacts:
                since_s = since.strftime("%d-%b-%Y")
                typ, data = imap.search(None, "FROM", f'"{email_addr}"', "SINCE", since_s)
                ids = data[0].split() if typ == "OK" and data and data[0] else []
                if ids:
                    replied[email_addr.lower()] = self._msg_date(imap, ids[0]) or dt.date.today()
                    continue
                for daemon in ("mailer-daemon", "postmaster"):
                    typ, data = imap.search(None, "FROM", f'"{daemon}"', "SINCE", since_s,
                                            "BODY", f'"{email_addr}"')
                    if typ == "OK" and data and data[0]:
                        bounced.add(email_addr.lower())
                        break
        finally:
            try:
                imap.logout()
            except Exception:
                pass
        return replied, bounced

    @staticmethod
    def _msg_date(imap, msg_id):
        typ, data = imap.fetch(msg_id, "(INTERNALDATE)")
        if typ == "OK" and data and data[0]:
            m = re.search(rb'INTERNALDATE "([^"]+)"', data[0] if isinstance(data[0], bytes) else data[0][0])
            if m:
                return dt.datetime.strptime(m.group(1).decode()[:11], "%d-%b-%Y").date()
        return None

    def close(self):
        if self.smtp is not None:
            try:
                self.smtp.quit()
            except Exception:
                pass


APPLESCRIPT_SEND = r"""
on run argv
    set toAddr to item 1 of argv
    set subj to item 2 of argv
    set bodyText to item 3 of argv
    set fromAddr to item 4 of argv
    with timeout of 300 seconds
        tell application "Mail"
            set m to make new outgoing message with properties {subject:subj, content:bodyText, visible:false}
            if fromAddr is not "" then set sender of m to fromAddr
            tell m to make new to recipient at end of to recipients with properties {address:toAddr}
            if (count of argv) > 4 then
                repeat with i from 5 to (count of argv)
                    set f to POSIX file (item i of argv)
                    tell content of m to make new attachment with properties {file name:f} at after last paragraph
                end repeat
                delay 3
            end if
            set ok to send m
        end tell
    end timeout
    if ok then return "OK"
    return "FAILED"
end run
"""

APPLESCRIPT_CHECK = r"""
on run argv
    set addr to item 1 of argv
    set d to current date
    set year of d to (item 2 of argv) as integer
    set day of d to 1
    set month of d to (item 3 of argv) as integer
    set day of d to (item 4 of argv) as integer
    set time of d to 0
    with timeout of 600 seconds
        tell application "Mail"
            set hits to (messages of inbox whose sender contains addr and date received >= d)
            if (count of hits) > 0 then
                set r to date received of item 1 of hits
                return "REPLY " & ((year of r) as text) & "-" & ((month of r as integer) as text) & "-" & ((day of r) as text)
            end if
            set ndrs to (messages of inbox whose date received >= d and (subject contains "Undeliverable" or subject contains "Delivery Status Notification" or subject contains "Mail delivery failed" or subject contains "Returned mail"))
            repeat with m in ndrs
                if content of m contains addr then return "BOUNCE"
            end repeat
        end tell
    end timeout
    return "NONE"
end run
"""


class AppleMailMailer:
    """Sends and checks replies through the Mac's built-in Mail app.

    Add your university Outlook / Microsoft 365 account to Mail once (Mail > Settings > Accounts,
    'Microsoft Exchange' or 'Sign in with Microsoft'). No password is stored by this script."""

    def __init__(self, cfg):
        if sys.platform != "darwin":
            sys.exit("method = applemail only works on a Mac.")
        s = cfg["email"]
        address = s.get("address", "").strip()
        name = s.get("display_name", "").strip()
        self.sender = f"{name} <{address}>" if name and address else address

    @staticmethod
    def _osascript(script, *args, timeout=900):
        import subprocess
        r = subprocess.run(["osascript", "-e", script, *[str(a) for a in args]],
                           capture_output=True, text=True, timeout=timeout)
        if r.returncode != 0:
            raise RuntimeError((r.stderr or r.stdout).strip() or "osascript failed")
        return r.stdout.strip()

    def check_connection(self):
        self._osascript('tell application "Mail" to return (count of accounts) as text')

    def send(self, to, subject, body, in_reply_to=None, attachments=()):
        out = self._osascript(APPLESCRIPT_SEND, to, subject, body, self.sender,
                              *[str(Path(a).resolve()) for a in attachments])
        if out != "OK":
            raise RuntimeError("Mail app refused to send the message")
        return ""

    def find_replies_and_bounces(self, contacts):
        replied, bounced = {}, set()
        for email_addr, since in contacts:
            out = self._osascript(APPLESCRIPT_CHECK, email_addr, since.year, since.month, since.day)
            if out.startswith("REPLY"):
                y, m, d = (int(x) for x in out.split()[1].split("-"))
                replied[email_addr.lower()] = dt.date(y, m, d)
            elif out == "BOUNCE":
                bounced.add(email_addr.lower())
        return replied, bounced

    def close(self):
        pass


def make_mailer(cfg):
    method = cfg["email"].get("method", "applemail").strip().lower()
    return SmtpImapMailer(cfg) if method == "smtp" else AppleMailMailer(cfg)


# --------------------------------------------------------------------------- core

def missing_fields(row, needed):
    return [c for c in needed if not str(row.get(c) or "").strip()]


def plan(sheet, cfg, today):
    """Decide what to send. Returns (initial_rows, followup_rows)."""
    s = cfg["settings"]
    follow_after = s.getint("follow_up_after_days", 7)
    initial, followups = [], []
    for row in sheet.rows():
        status = str(row.get(COL_STATUS) or "").strip()
        if status.lower() in STATUS_SKIP:
            continue
        if not status:
            initial.append(row)
        elif status == STATUS_SENT and s.getboolean("send_follow_ups", True):
            replied = str(row.get(COL_REPLIED) or "").strip().lower()
            sent_on = as_date(row.get(COL_DATE_SENT))
            if replied.startswith("y") or row.get(COL_FOLLOW_UP) or not sent_on:
                continue
            if (today - sent_on).days >= follow_after:
                followups.append(row)
    return initial, followups


def cmd_preview(cfg, sheet_path):
    today = dt.date.today()
    sheet = Sheet(sheet_path)
    templates = Templates(cfg)
    initial, followups = plan(sheet, cfg, today)
    print(f"=== {len(initial)} new contact(s), {len(followups)} follow-up(s) due ===\n")
    for row in initial:
        tpl = templates.for_row(row)
        if not tpl:
            print(f"--- row {row['_row']}: SKIPPED, no template for Commonality "
                  f"'{row.get(COL_COMMON) or ''}' (have: {templates.names()})\n")
            continue
        if tpl.missing_attachment:
            print(f"--- row {row['_row']}: SKIPPED, email says 'attached' but no resume PDF in templates/{tpl.name}/\n")
            continue
        miss = missing_fields(row, tpl.needed)
        if miss:
            print(f"--- row {row['_row']}: SKIPPED, missing {', '.join(miss)}\n")
            continue
        subj, body = tpl.first(row)
        att = f"\nAttachments: {', '.join(a.name for a in tpl.attachments)}" if tpl.attachments else ""
        print(f"--- row {row['_row']}  To: {row[COL_EMAIL]}\nSubject: {subj}{att}\n\n{body}")
    for row in followups:
        tpl = templates.for_row(row)
        if tpl:
            subj, body = tpl.follow_up(row)
            print(f"--- FOLLOW-UP row {row['_row']}  To: {row[COL_EMAIL]}\nSubject: {subj}\n\n{body}")


def cmd_test(cfg):
    mailer = make_mailer(cfg)
    me = cfg["email"]["address"].strip()
    for tpl in Templates(cfg).sets.values():
        if tpl.missing_attachment:
            print(f"Skipping '{tpl.name}': the email mentions an attachment but templates/{tpl.name}/ has no PDF.")
            continue
        sample = {COL_FIRST: "Alex", COL_LAST: "Sample", COL_EMAIL: me, COL_COMPANY: "Acme Corp",
                  COL_ROLE: "Account Executive", COL_AREA: "sales", COL_COMMON: tpl.name}
        subj, body = tpl.first(sample)
        sample[COL_SUBJECT] = "[TEST] " + subj
        mid = mailer.send(me, sample[COL_SUBJECT], body, attachments=tpl.attachments)
        subj2, body2 = tpl.follow_up(sample)
        mailer.send(me, subj2, body2, in_reply_to=mid)
        print(f"Sent test first email + follow-up for template '{tpl.name}' to {me}.")
    mailer.close()


def cmd_run(cfg, sheet_path, force=False):
    today = dt.date.today()
    if today.weekday() >= 5 and not force:
        log.info("Weekend — nothing sent. (Use --force to override.)")
        return
    s = cfg["settings"]
    max_per_day = s.getint("max_emails_per_day", 25)
    delay_min, delay_max = s.getint("min_seconds_between_emails", 30), s.getint("max_seconds_between_emails", 90)

    templates = Templates(cfg)

    # Back up the sheet before touching it.
    backup_dir = HERE / "backups"
    backup_dir.mkdir(exist_ok=True)
    shutil.copy2(sheet_path, backup_dir / f"{Path(sheet_path).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")

    sheet = Sheet(sheet_path)
    mailer = make_mailer(cfg)

    # 1) Update Replied / Delivery from the inbox before deciding on follow-ups.
    watch = []
    for row in sheet.rows():
        sent_on = as_date(row.get(COL_DATE_SENT))
        if row.get(COL_STATUS) in (STATUS_SENT, STATUS_FOLLOWED_UP) and sent_on and row.get(COL_EMAIL):
            watch.append((str(row[COL_EMAIL]).strip(), sent_on))
    try:
        replied, bounced = mailer.find_replies_and_bounces(watch)
    except Exception as e:
        log.warning("Couldn't check inbox for replies (%s). Follow-ups skipped today to be safe.", e)
        replied, bounced, s["send_follow_ups"] = {}, set(), "false"
    for row in sheet.rows():
        e = str(row.get(COL_EMAIL) or "").strip().lower()
        if row.get(COL_STATUS) not in (STATUS_SENT, STATUS_FOLLOWED_UP):
            continue
        if e in replied:
            sheet.set(row, COL_REPLIED, "Yes")
            sheet.set(row, COL_STATUS, STATUS_REPLIED)
            log.info("Reply detected from %s", e)
        elif e in bounced:
            sheet.set(row, COL_DELIVERY, "Bounced")
            sheet.set(row, COL_STATUS, STATUS_BOUNCED)
            log.info("Bounce detected for %s", e)
        elif not row.get(COL_REPLIED):
            sheet.set(row, COL_REPLIED, "No")
    sheet.save()

    # 2) Send.
    initial, followups = plan(sheet, cfg, today)
    sent = 0

    def pause():
        if delay_max > 0:
            time.sleep(random.uniform(delay_min, delay_max))

    try:
        for row in followups:
            if sent >= max_per_day:
                break
            to = str(row[COL_EMAIL]).strip()
            tpl = templates.for_row(row)
            if not tpl:
                log.warning("No template for %s's Commonality '%s'; follow-up not sent.", to, row.get(COL_COMMON))
                continue
            subj, body = tpl.follow_up(row)
            try:
                if sent:
                    pause()
                mailer.send(to, subj, body, in_reply_to=row.get(COL_MSG_ID) or None)
                sheet.set(row, COL_FOLLOW_UP, today, "yyyy-mm-dd")
                sheet.set(row, COL_STATUS, STATUS_FOLLOWED_UP)
                sent += 1
                log.info("Follow-up sent to %s", to)
            except Exception as e:
                sheet.set(row, COL_FOLLOW_UP, f"Failed: {e}"[:200])
                log.error("Follow-up to %s failed: %s", to, e)
            sheet.save()

        for row in initial:
            if sent >= max_per_day:
                log.info("Daily limit (%d) reached; the rest go out on the next workday.", max_per_day)
                break
            tpl = templates.for_row(row)
            if not tpl:
                sheet.set(row, COL_DELIVERY, f"Waiting: no template for '{row.get(COL_COMMON) or ''}'"
                          f" (have: {templates.names()})")
                continue
            if tpl.missing_attachment:
                sheet.set(row, COL_DELIVERY, f"Waiting: put your resume PDF in templates/{tpl.name}/")
                continue
            miss = missing_fields(row, tpl.needed)
            if miss:
                sheet.set(row, COL_DELIVERY, f"Waiting: fill in {', '.join(miss)}")
                continue
            to = str(row[COL_EMAIL]).strip()
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", to):
                sheet.set(row, COL_DELIVERY, "Invalid email address")
                continue
            subj, body = tpl.first(row)
            try:
                if sent:
                    pause()
                msg_id = mailer.send(to, subj, body, attachments=tpl.attachments)
                sheet.set(row, COL_STATUS, STATUS_SENT)
                sheet.set(row, COL_DATE_SENT, today, "yyyy-mm-dd")
                sheet.set(row, COL_DELIVERY, "Sent")
                sheet.set(row, COL_REPLIED, "No")
                sheet.set(row, COL_MSG_ID, msg_id)
                sheet.set(row, COL_SUBJECT, subj)
                sent += 1
                log.info("First email sent to %s", to)
            except Exception as e:
                sheet.set(row, COL_STATUS, STATUS_FAILED)
                sheet.set(row, COL_DELIVERY, f"Failed: {e}"[:200])
                log.error("Email to %s failed: %s", to, e)
            sheet.save()
    finally:
        sheet.save()
        mailer.close()
    log.info("Done. %d email(s) sent today.", sent)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("command", choices=["init", "preview", "test", "run", "check"])
    p.add_argument("--force", action="store_true", help="run even on weekends")
    args = p.parse_args()
    setup_logging()

    if args.command == "init":
        target = HERE / "outreach.xlsx"
        if target.exists():
            sys.exit("outreach.xlsx already exists — not overwriting it.")
        shutil.copy2(TEMPLATE_XLSX, target)
        print(f"Created {target}")
        return

    cfg = load_config()
    sheet_path = HERE / cfg["settings"].get("excel_file", "outreach.xlsx")
    if args.command in ("preview", "run") and not sheet_path.exists():
        sys.exit(f"{sheet_path.name} not found. Run:  python outreach.py init")

    if args.command == "preview":
        cmd_preview(cfg, sheet_path)
    elif args.command == "check":
        make_mailer(cfg).check_connection()
        print("Login OK.")
    elif args.command == "test":
        cmd_test(cfg)
    else:
        try:
            cmd_run(cfg, sheet_path, force=args.force)
        except Exception:
            log.exception("Run crashed")
            raise


if __name__ == "__main__":
    main()
